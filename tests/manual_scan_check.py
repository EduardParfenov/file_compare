"""Ручная проверка полного сценария на сканах (задача 8.4).

Приложение поднимается в-process с подменённой моделью: настоящий эндпоинт
требует авторизации (отвечает 401), поэтому проверяется всё, кроме качества
чтения реальной vision-моделью. Сценарий проходит по тем же HTTP-эндпоинтам,
что и браузер, и проверяет контракт, который потребляет app.js.

Запуск: python -m tests.manual_scan_check
"""

import hashlib
import sys
from pathlib import Path
from types import SimpleNamespace

sys.path.insert(0, str(Path(__file__).parent))
from pdf_fixtures import write_scan_pdf

from app import create_app
from app.services import conversion, jobs, vision_ocr

FAILURES: list[str] = []


def check(condition: bool, message: str) -> None:
    status = "ok  " if condition else "FAIL"
    print(f"  [{status}] {message}")
    if not condition:
        FAILURES.append(message)


class StubChat:
    """Модель: классифицирует фрагменты и читает страницы как OCR.

    Текст страницы выводится из отпечатка изображения, а не из счётчика
    вызовов: тогда одинаковые страницы дают одинаковый текст, а отличающиеся
    — разный. Именно это и требуется от настоящей OCR-модели, и без такой
    связи проверка на свёрнутые диапазоны проходила бы вхолостую.
    """

    def invoke(self, messages):
        if isinstance(messages[1].content, list):  # мультимодальный запрос
            return SimpleNamespace(content=self._page_text(messages))
        return SimpleNamespace(content='{"label": "changed"}')

    @staticmethod
    def _page_text(messages) -> str:
        url = next(
            part["image_url"]["url"]
            for part in messages[1].content
            if part["type"] == "image_url"
        )
        digest = hashlib.sha1(url.encode()).hexdigest()[:6]
        return (
            f"# Страница\n\n"
            f"Абзац страницы содержит текст для сравнения.\n\n"
            "| колонка | значение |\n| --- | --- |\n"
            f"| строка | {digest} |"
        )


def build(tmp: Path):
    app = create_app(
        {
            "TESTING": True,
            "UPLOAD_DIR": str(tmp / "uploads"),
            "ALLOWED_EXTENSIONS": {".pdf", ".docx"},
            "JOBS_SYNCHRONOUS": True,
            "LLM_MODEL": "stub-vl",
            "PDF_RENDER_DPI": "150",
        }
    )
    return app.test_client()


def upload(client, path: Path) -> str:
    with open(path, "rb") as handle:
        response = client.post(
            "/api/upload",
            data={"file": (handle, path.name)},
            content_type="multipart/form-data",
        )
    assert response.status_code == 200, response.get_json()
    return response.get_json()["upload_id"]


def main() -> int:
    tmp = Path("/tmp/opencode/manual-scan-check")
    tmp.mkdir(parents=True, exist_ok=True)

    # Файл 1: четыре страницы. Файл 2: те же страницы, но вторая изменена.
    file1 = write_scan_pdf(
        tmp / "contract_v1.pdf",
        [
            [(40, 320, "Contract page one")],
            [(40, 320, "Contract page two")],
            [(40, 320, "Contract page three")],
            [(40, 320, "Contract page four")],
        ],
    )
    file2 = write_scan_pdf(
        tmp / "contract_v2.pdf",
        [
            [(40, 320, "Contract page one")],
            [(40, 320, "Contract page two CHANGED")],
            [(40, 320, "Contract page three")],
            [(40, 320, "Contract page four")],
        ],
    )

    vision_ocr.clear_cache()
    client = build(tmp)
    # Чтение страниц подменяется заглушкой: настоящий эндпоинт требует ключа.
    # Сбой модели проверяется отдельно, через подмену заглушки ниже.
    original_reader = conversion._reader_for
    healthy_chat = StubChat()
    jobs._JOBS.clear()
    conversion._reader_for = lambda chat: original_reader(healthy_chat)

    print("\n1. Загрузка и сравнение сканов")
    upload_id_1 = upload(client, Path(file1))
    upload_id_2 = upload(client, Path(file2))
    response = client.post(
        "/api/compare",
        json={"upload_id_1": upload_id_1, "upload_id_2": upload_id_2},
    )
    check(response.status_code == 202, "POST /api/compare принят (202)")
    job_id = response.get_json()["job_id"]

    print("\n2. Статус задачи")
    status = client.get(f"/api/jobs/{job_id}").get_json()
    check(status["status"] == "done", f"задача завершена ({status['status']})")
    check(
        status["stage"] in {"converting", "diffing", "llm"},
        f"ключ этапа прежний ({status['stage']})",
    )

    result = status["result"]
    rows = result["rows"]
    pages = result["pages"]

    print("\n3. Контракт результата для клиента")
    check("pages" in result, "результат содержит сведения о страницах")
    check("crops" in result, "результат содержит кропы")
    check("crops_truncated" in result, "результат содержит признак усечения кропов")
    check(
        pages["page_count"] == {"left": 4, "right": 4},
        f"число страниц известно ({pages['page_count']})",
    )

    unchanged = {tuple(pair) for pair in pages["unchanged"]}
    check((1, 1) in unchanged, "неизменённая страница 1 помечена совпавшей")
    check((2, 2) not in unchanged, "изменённая страница 2 не помечена совпавшей")
    check(
        (3, 3) in unchanged and (4, 4) in unchanged, "страницы 3–4 помечены совпавшими"
    )

    collapsed = [
        row["left"]["collapsed"]
        for row in rows
        if row.get("left") and row["left"].get("collapsed")
    ]
    check(
        all(0 < span[0] <= span[1] <= 4 for span in collapsed),
        f"свёрнутые диапазоны корректны ({collapsed})",
    )

    page_rows = [row for row in rows if row.get("left") and row["left"].get("page")]
    check(bool(page_rows), "строки содержат номер страницы")
    check(
        all(
            {"text", "change", "html", "images"} <= set(row["left"])
            for row in page_rows
        ),
        "стороны строк сохраняют прежний контракт (text/change/html/images)",
    )

    # Подсветка живёт в segments для абзацев и в cell_segments для строк
    # таблиц — обе формы должны дойти до клиента.
    paragraph_highlights = [
        row for row in rows if (row.get("left") or {}).get("segments")
    ]
    cell_highlights = [
        row for row in rows if (row.get("left") or {}).get("cell_segments")
    ]
    check(
        bool(paragraph_highlights or cell_highlights),
        "пословная подсветка присутствует в изменённых строках",
    )
    check(bool(cell_highlights), "строки таблиц подсвечены по ячейкам")

    print("\n4. Кропы визуально изменённых областей")
    crop_pages = sorted(int(page) for page in result["crops"] if result["crops"][page])
    check(crop_pages == [2], f"кропы есть для изменённой страницы ({crop_pages})")
    crop = result["crops"]["2"][0] if result["crops"].get("2") else ""
    check(crop.startswith("data:image/jpeg;base64,"), "кроп передан как data-URI")
    check(
        crop_pages == [2] and 1 not in result["crops"],
        "у неизменённых страниц кропов нет",
    )

    print("\n5. Раскрытие свёрнутого диапазона")
    range_row = next(
        (row for row in rows if row.get("right") and row["right"].get("collapsed")),
        None,
    )
    check(range_row is not None, "в результате есть строка свёрнутого диапазона")
    first, last = range_row["right"]["collapsed"]
    expansion = client.post(
        f"/api/jobs/{job_id}/pages",
        json={"side": "right", "first": first, "last": last},
    )
    check(expansion.status_code == 200, "POST /api/jobs/<id>/pages принят")
    body = expansion.get_json()
    check(
        body["pages"] != [] and all("page" in block for block in body["pages"]),
        f"диапазон {first}–{last} развёрнут в блоки с номерами страниц",
    )

    bad = client.post(
        f"/api/jobs/{job_id}/pages", json={"side": "right", "first": 9, "last": 9}
    )
    check(bad.status_code == 409, "страница вне документа отклоняется (409)")

    print("\n6. Полное совпадение документов")
    same = write_scan_pdf(
        tmp / "identical.pdf",
        [[(40, 320, "Contract page one")], [(40, 320, "Contract page two")]],
    )
    upload_id_3 = upload(client, Path(same))
    identical_job = client.post(
        "/api/compare", json={"upload_id_1": upload_id_3, "upload_id_2": upload_id_3}
    ).get_json()["job_id"]
    identical = client.get(f"/api/jobs/{identical_job}").get_json()
    check(identical["status"] == "done", "идентичные документы обработаны")
    same_pages = identical["result"]["pages"]
    check(
        len(same_pages["unchanged"]) == 2,
        "обе страницы признаны совпавшими без обращения к модели",
    )

    print("\n7. Совместимость с .docx")
    from docx import Document

    document = Document()
    document.add_paragraph("Сумма договора составляет 100 000 рублей")
    docx_path = tmp / "plain.docx"
    document.save(str(docx_path))
    changed_docx = tmp / "plain2.docx"
    document2 = Document()
    document2.add_paragraph("Сумма договора составляет 120 000 рублей")
    document2.save(str(changed_docx))

    docx_job = client.post(
        "/api/compare",
        json={
            "upload_id_1": upload(client, docx_path),
            "upload_id_2": upload(client, changed_docx),
        },
    ).get_json()["job_id"]
    docx_result = client.get(f"/api/jobs/{docx_job}").get_json()
    check(docx_result["status"] == "done", ".docx обработан прежним путём")
    check("pages" not in docx_result["result"], ".docx не получает страниц и свёртки")
    docx_rows = docx_result["result"]["rows"]
    highlighted = [
        row
        for row in docx_rows
        if (row.get("left") or {}).get("segments")
        or (row.get("left") or {}).get("cell_segments")
    ]
    check(bool(highlighted), "пословная подсветка .docx не сломалась")
    changed_words = [
        seg["text"]
        for row in highlighted
        for seg in (row["left"].get("segments") or [])
        if seg["type"] in {"del", "add"}
    ]
    check(
        any("100" in word or "120" in word for word in changed_words),
        f"подсвечены именно изменённые числа ({changed_words})",
    )
    check(
        all(
            not (row.get("left") or {}).get("collapsed")
            and not (row.get("left") or {}).get("page")
            for row in docx_rows
        ),
        "строки .docx не получили страниц и свёрток",
    )

    print("\n8. Деградация при недоступной модели")
    vision_ocr.clear_cache()

    class BrokenChat(StubChat):
        def invoke(self, messages):
            if isinstance(messages[1].content, list):
                raise ConnectionError("model down")
            return SimpleNamespace(content='{"label": "changed"}')

    conversion._reader_for = lambda chat: original_reader(BrokenChat())
    degraded_job = client.post(
        "/api/compare",
        json={"upload_id_1": upload_id_1, "upload_id_2": upload_id_2},
    ).get_json()["job_id"]
    degraded = client.get(f"/api/jobs/{degraded_job}").get_json()
    check(
        degraded["status"] == "done",
        f"сбой модели не роняет задачу ({degraded['status']}: {degraded.get('error')})",
    )
    if degraded["status"] != "done":
        raise SystemExit(1)
    check(
        degraded["result"]["semantic"] is False,
        "признак деградации поднят в результате",
    )
    unreadable_pages = sorted(
        {
            row["left"]["page"]
            for row in degraded["result"]["rows"]
            if row.get("left") and row["left"].get("unreadable")
        }
    )
    check(
        unreadable_pages == [2],
        f"непрочитанная страница показана, а не свёрнута ({unreadable_pages})",
    )
    conversion._reader_for = original_reader

    print()
    if FAILURES:
        print(f"ПРОВАЛЕНО проверок: {len(FAILURES)}")
        for message in FAILURES:
            print(f"  - {message}")
        return 1
    print("Все проверки пройдены")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
